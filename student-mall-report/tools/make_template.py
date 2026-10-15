"""
Jana templat laporan kewangan (1 helaian Excel) dan benamkannya dalam Index.html.

Gunakan:  python3 tools/make_template.py
Hasil:    assets/Templat_Laporan_Kewangan.xlsx  dan  kemas kini TEMPLATE_B64 dalam Index.html
Keperluan: pip install openpyxl
"""
import base64, os, re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'assets', 'Templat_Laporan_Kewangan.xlsx')

FIRST, LAST = 15, 58          # baris jadual transaksi (44 baris, muat satu halaman A4)
TOTAL = LAST + 1

wb = Workbook()
ws = wb.active
ws.title = 'Laporan'

thin = Side(style='thin', color='000000')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
peach = PatternFill('solid', fgColor='F8CBAD')
grey = PatternFill('solid', fgColor='C0C0C0')
input_fill = PatternFill('solid', fgColor='FFF2CC')
bold = Font(name='Arial', bold=True, size=10)
base = Font(name='Arial', size=10)
note = Font(name='Arial', italic=True, size=9, color='595959')

for col, w in zip('ABCD', (25, 48, 20, 20)):
    ws.column_dimensions[col].width = w

# Tajuk
ws['A1'] = 'LAPORAN KEWANGAN BULANAN'
ws['A1'].font = Font(name='Arial', bold=True, size=14)
ws.merge_cells('A1:D1')
ws['A2'] = 'Isi sel berwarna kuning dan jadual di bawah. Jumlah dan untung/rugi dikira secara automatik.'
ws['A2'].font = note
ws.merge_cells('A2:D2')

# Maklumat laporan
for r, label in ((4, 'Nama perniagaan'), (5, 'Lot sewaan'), (6, 'Bulan dan tahun laporan'), (7, 'Disediakan oleh')):
    ws.cell(r, 1, label).font = bold
    c = ws.cell(r, 2)
    c.fill = input_fill; c.border = box; c.font = base
ws['C5'] = 'Contoh: Student Mall - Lot 3'
ws['C6'] = 'Contoh: Mei 2026'
for ref in ('C5', 'C6'):
    ws[ref].font = Font(name='Arial', size=9, color='7F7F7F', italic=True)

# Ringkasan
ws.merge_cells('A9:D9')
ws['A9'] = 'RINGKASAN'
ws['A9'].font = bold; ws['A9'].alignment = Alignment(horizontal='center')
for col in 'ABCD':
    ws[f'{col}9'].fill = grey
rows = ((10, 'A', 'Jumlah Pendapatan', f'=C{TOTAL}'),
        (11, 'B', 'Jumlah Perbelanjaan', f'=D{TOTAL}'),
        (12, '', 'Untung / Rugi (A - B)', '=C10-C11'))
for r, a, b, f in rows:
    ws.cell(r, 1, a).font = base
    ws.cell(r, 2, b).font = bold
    c = ws.cell(r, 3, f); c.font = bold; c.number_format = '#,##0.00;[Red]-#,##0.00'

# Jadual transaksi
header_row = 14
for i, h in enumerate(('Tarikh', 'Perkara', 'Pendapatan (RM)', 'Perbelanjaan (RM)'), start=1):
    c = ws.cell(header_row, i, h)
    c.font = Font(name='Times New Roman', bold=True, size=11)
    c.fill = peach; c.border = box
    c.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[header_row].height = 22

samples = ('Baki dibawa dari bulan lepas', 'Jualan', 'Bayaran sewa', 'Bahan mentah', 'Gaji')
for r in range(FIRST, LAST + 1):
    for col in range(1, 5):
        c = ws.cell(r, col); c.border = box; c.font = base
    ws.cell(r, 1).number_format = 'dd mmm yyyy'
    ws.cell(r, 1).alignment = Alignment(horizontal='right')
    ws.cell(r, 3).number_format = '#,##0.00'
    ws.cell(r, 4).number_format = '#,##0.00'
for i, text in enumerate(samples):
    ws.cell(FIRST + i, 2, text)

# Jumlah
ws.cell(TOTAL, 2, 'Jumlah').font = bold
ws.cell(TOTAL, 3, f'=SUM(C{FIRST}:C{LAST})')
ws.cell(TOTAL, 4, f'=SUM(D{FIRST}:D{LAST})')
for col in range(1, 5):
    c = ws.cell(TOTAL, col); c.border = box; c.font = bold
    if col >= 3: c.number_format = '#,##0.00'

ws.cell(TOTAL + 2, 1, 'Padam atau tukar baris contoh jika tidak berkaitan. Selepas siap, simpan atau cetak sebagai PDF dan muat naik ke sistem.').font = note
ws.merge_cells(start_row=TOTAL + 2, start_column=1, end_row=TOTAL + 2, end_column=4)

# Cetak: A4, muat 1 lebar, ulang baris tajuk jadual
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.orientation = 'portrait'
ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.print_title_rows = f'{header_row}:{header_row}'
ws.print_area = f'A1:D{TOTAL + 2}'
ws.page_margins.left = ws.page_margins.right = 0.5
ws.sheet_view.showGridLines = False

wb.save(OUT)
print('Disimpan:', OUT, os.path.getsize(OUT), 'bait')

# Benamkan dalam Index.html
b64 = base64.b64encode(open(OUT, 'rb').read()).decode()
idx = os.path.join(ROOT, 'Index.html')
html = open(idx, encoding='utf-8').read()
new = re.sub(r"(var TEMPLATE_B64 = ')[^']*(';)", lambda m: m.group(1) + b64 + m.group(2), html, count=1)
if new == html and 'var TEMPLATE_B64' not in html:
    print('Index.html belum mempunyai TEMPLATE_B64; tambah dahulu.')
else:
    open(idx, 'w', encoding='utf-8').write(new)
    print('Index.html dikemas kini')
